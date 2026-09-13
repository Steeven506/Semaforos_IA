from fastapi import FastAPI, Query
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse
import psycopg2
from datetime import datetime
from typing import Optional
import csv
import io
import tempfile
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

app = FastAPI(title="Traffic Detection API", version="1.0")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

def get_db_connection():
    return psycopg2.connect(
        host='localhost',
        port='5432',
        database='traffic_db',
        user='postgres',
        password='postgres'
    )

@app.get("/")
def root():
    return {
        "message": "Traffic Detection API",
        "version": "1.0",
        "endpoints": [
            "/stats/today",
            "/stats/hourly",
            "/stats/weekly",
            "/stats/cameras",
            "/stats/latest",
            "/stats/summary",
            "/export/csv?tipo=diario",
            "/export/excel?tipo=diario"
        ]
    }

@app.get("/stats/today")
def get_today_stats():
    conn = get_db_connection()
    cursor = conn.cursor()
    today = datetime.now().strftime('%Y-%m-%d')
    
    cursor.execute('''
        SELECT 
            COALESCE(SUM(personas), 0) as total_personas,
            COALESCE(SUM(carros), 0) as total_carros,
            COALESCE(SUM(motos), 0) as total_motos,
            COALESCE(SUM(buses), 0) as total_buses,
            COALESCE(SUM(camiones), 0) as total_camiones,
            COALESCE(SUM(total), 0) as total_vehiculos,
            COUNT(*) as total_registros
        FROM detections_log
        WHERE DATE(timestamp) = %s
    ''', (today,))
    
    result = cursor.fetchone()
    cursor.close()
    conn.close()
    
    if result:
        return {
            "total_personas": result[0] or 0,
            "total_carros": result[1] or 0,
            "total_motos": result[2] or 0,
            "total_buses": result[3] or 0,
            "total_camiones": result[4] or 0,
            "total_vehiculos": result[5] or 0,
            "total_registros": result[6] or 0
        }
    return {"total_personas": 0, "total_carros": 0, "total_motos": 0, "total_buses": 0, "total_camiones": 0, "total_vehiculos": 0, "total_registros": 0}

@app.get("/stats/hourly")
def get_hourly_stats():
    conn = get_db_connection()
    cursor = conn.cursor()
    today = datetime.now().strftime('%Y-%m-%d')
    
    cursor.execute('''
        SELECT 
            TO_CHAR(timestamp, 'HH24:00') as hour,
            COALESCE(SUM(personas), 0) as personas,
            COALESCE(SUM(carros), 0) as carros,
            COALESCE(SUM(motos), 0) as motos,
            COALESCE(SUM(total), 0) as total
        FROM detections_log
        WHERE DATE(timestamp) = %s
        GROUP BY hour
        ORDER BY hour
    ''', (today,))
    
    results = cursor.fetchall()
    cursor.close()
    conn.close()
    
    data = []
    for row in results:
        data.append({
            "hour": row[0],
            "personas": row[1] or 0,
            "carros": row[2] or 0,
            "motos": row[3] or 0,
            "total": row[4] or 0
        })
    return data

@app.get("/stats/weekly")
def get_weekly_stats():
    conn = get_db_connection()
    cursor = conn.cursor()
    
    cursor.execute('''
        SELECT 
            DATE(timestamp) as date,
            COALESCE(SUM(personas), 0) as personas,
            COALESCE(SUM(carros), 0) as carros,
            COALESCE(SUM(motos), 0) as motos,
            COALESCE(SUM(total), 0) as total
        FROM detections_log
        WHERE DATE(timestamp) >= CURRENT_DATE - INTERVAL '7 days'
        GROUP BY DATE(timestamp)
        ORDER BY date
    ''')
    
    results = cursor.fetchall()
    cursor.close()
    conn.close()
    
    data = []
    for row in results:
        data.append({
            "date": row[0].strftime('%Y-%m-%d') if row[0] else "",
            "personas": row[1] or 0,
            "carros": row[2] or 0,
            "motos": row[3] or 0,
            "total": row[4] or 0
        })
    return data

@app.get("/stats/latest")
def get_latest_detections(limit: Optional[int] = 20):
    conn = get_db_connection()
    cursor = conn.cursor()
    
    cursor.execute('''
        SELECT 
            timestamp,
            personas,
            carros,
            motos,
            buses,
            camiones,
            total
        FROM detections_log
        ORDER BY timestamp DESC
        LIMIT %s
    ''', (limit,))
    
    results = cursor.fetchall()
    cursor.close()
    conn.close()
    
    data = []
    for row in results:
        data.append({
            "timestamp": row[0].strftime('%Y-%m-%d %H:%M:%S') if row[0] else "",
            "personas": row[1] or 0,
            "carros": row[2] or 0,
            "motos": row[3] or 0,
            "buses": row[4] or 0,
            "camiones": row[5] or 0,
            "total": row[6] or 0
        })
    return data

@app.get("/stats/cameras")
def get_cameras():
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute('SELECT id, name, url, camera_type FROM cameras')
    results = cursor.fetchall()
    cursor.close()
    conn.close()
    
    data = []
    for row in results:
        data.append({
            "id": row[0],
            "name": row[1],
            "url": row[2],
            "camera_type": row[3]
        })
    return data

@app.get("/stats/summary")
def get_summary():
    conn = get_db_connection()
    cursor = conn.cursor()
    today = datetime.now().strftime('%Y-%m-%d')
    
    cursor.execute('''
        SELECT 
            COALESCE(SUM(total), 0) as total_vehiculos,
            COALESCE(AVG(total), 0) as promedio_por_minuto,
            COALESCE(MAX(total), 0) as maximo_por_minuto,
            COUNT(*) as total_minutos
        FROM detections_log
        WHERE DATE(timestamp) = %s
    ''', (today,))
    
    result = cursor.fetchone()
    cursor.close()
    conn.close()
    
    return {
        "total_vehiculos": result[0] or 0,
        "promedio_por_minuto": round(result[1], 2) if result[1] else 0,
        "maximo_por_minuto": result[2] or 0,
        "total_minutos": result[3] or 0
    }

@app.get("/export/csv")
def export_csv(
    tipo: str = Query("diario", description="diario, semanal o mensual")
):
    conn = get_db_connection()
    cursor = conn.cursor()
    
    if tipo == "diario":
        fecha = datetime.now().strftime('%Y-%m-%d')
        cursor.execute('''
            SELECT timestamp, personas, carros, motos, buses, camiones, total
            FROM detections_log
            WHERE DATE(timestamp) = %s
            ORDER BY timestamp
        ''', (fecha,))
        nombre = f"reporte_diario_{fecha}.csv"
        headers = ['Timestamp', 'Personas', 'Carros', 'Motos', 'Buses', 'Camiones', 'Total']
    
    elif tipo == "semanal":
        cursor.execute('''
            SELECT DATE(timestamp) as fecha, 
                   COALESCE(SUM(personas), 0) as personas,
                   COALESCE(SUM(carros), 0) as carros,
                   COALESCE(SUM(motos), 0) as motos,
                   COALESCE(SUM(buses), 0) as buses,
                   COALESCE(SUM(camiones), 0) as camiones,
                   COALESCE(SUM(total), 0) as total
            FROM detections_log
            WHERE DATE(timestamp) >= CURRENT_DATE - INTERVAL '7 days'
            GROUP BY DATE(timestamp)
            ORDER BY fecha
        ''')
        nombre = f"reporte_semanal_{datetime.now().strftime('%Y-%m-%d')}.csv"
        headers = ['Fecha', 'Personas', 'Carros', 'Motos', 'Buses', 'Camiones', 'Total']
    
    elif tipo == "mensual":
        cursor.execute('''
            SELECT DATE(timestamp) as fecha, 
                   COALESCE(SUM(personas), 0) as personas,
                   COALESCE(SUM(carros), 0) as carros,
                   COALESCE(SUM(motos), 0) as motos,
                   COALESCE(SUM(buses), 0) as buses,
                   COALESCE(SUM(camiones), 0) as camiones,
                   COALESCE(SUM(total), 0) as total
            FROM detections_log
            WHERE DATE(timestamp) >= CURRENT_DATE - INTERVAL '30 days'
            GROUP BY DATE(timestamp)
            ORDER BY fecha
        ''')
        nombre = f"reporte_mensual_{datetime.now().strftime('%Y-%m-%d')}.csv"
        headers = ['Fecha', 'Personas', 'Carros', 'Motos', 'Buses', 'Camiones', 'Total']
    
    else:
        cursor.close()
        conn.close()
        return {"error": "Tipo no valido. Usa: diario, semanal o mensual"}
    
    results = cursor.fetchall()
    cursor.close()
    conn.close()
    
    if not results:
        return {"error": "No hay datos para exportar"}
    
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(headers)
    
    for row in results:
        writer.writerow(row)
    
    output.seek(0)
    
    with tempfile.NamedTemporaryFile(mode='w', delete=False, suffix='.csv', encoding='utf-8') as tmp:
        tmp.write(output.getvalue())
        tmp_path = tmp.name
    
    return FileResponse(
        path=tmp_path,
        filename=nombre,
        media_type='text/csv'
    )

@app.get("/export/excel")
def export_excel(
    tipo: str = Query("diario", description="diario, semanal o mensual")
):
    conn = get_db_connection()
    cursor = conn.cursor()
    
    if tipo == "diario":
        fecha = datetime.now().strftime('%Y-%m-%d')
        cursor.execute('''
            SELECT timestamp, personas, carros, motos, buses, camiones, total
            FROM detections_log
            WHERE DATE(timestamp) = %s
            ORDER BY timestamp
        ''', (fecha,))
        titulo = f"Reporte Diario - {fecha}"
        nombre = f"reporte_diario_{fecha}.xlsx"
        headers = ['Timestamp', 'Personas', 'Carros', 'Motos', 'Buses', 'Camiones', 'Total']
        
        cursor_resumen = conn.cursor()
        cursor_resumen.execute('''
            SELECT 
                COALESCE(SUM(personas), 0) as total_personas,
                COALESCE(SUM(carros), 0) as total_carros,
                COALESCE(SUM(motos), 0) as total_motos,
                COALESCE(SUM(buses), 0) as total_buses,
                COALESCE(SUM(camiones), 0) as total_camiones,
                COALESCE(SUM(total), 0) as total_vehiculos,
                COUNT(*) as total_registros,
                COALESCE(AVG(total), 0) as promedio
            FROM detections_log
            WHERE DATE(timestamp) = %s
        ''', (fecha,))
        resumen = cursor_resumen.fetchone()
        cursor_resumen.close()
        
    elif tipo == "semanal":
        cursor.execute('''
            SELECT DATE(timestamp) as fecha, 
                   COALESCE(SUM(personas), 0) as personas,
                   COALESCE(SUM(carros), 0) as carros,
                   COALESCE(SUM(motos), 0) as motos,
                   COALESCE(SUM(buses), 0) as buses,
                   COALESCE(SUM(camiones), 0) as camiones,
                   COALESCE(SUM(total), 0) as total
            FROM detections_log
            WHERE DATE(timestamp) >= CURRENT_DATE - INTERVAL '7 days'
            GROUP BY DATE(timestamp)
            ORDER BY fecha
        ''')
        titulo = f"Reporte Semanal - {datetime.now().strftime('%Y-%m-%d')}"
        nombre = f"reporte_semanal_{datetime.now().strftime('%Y-%m-%d')}.xlsx"
        headers = ['Fecha', 'Personas', 'Carros', 'Motos', 'Buses', 'Camiones', 'Total']
        
        cursor_resumen = conn.cursor()
        cursor_resumen.execute('''
            SELECT 
                COALESCE(SUM(personas), 0) as total_personas,
                COALESCE(SUM(carros), 0) as total_carros,
                COALESCE(SUM(motos), 0) as total_motos,
                COALESCE(SUM(buses), 0) as total_buses,
                COALESCE(SUM(camiones), 0) as total_camiones,
                COALESCE(SUM(total), 0) as total_vehiculos,
                COUNT(*) as total_registros,
                COALESCE(AVG(total), 0) as promedio
            FROM detections_log
            WHERE DATE(timestamp) >= CURRENT_DATE - INTERVAL '7 days'
        ''')
        resumen = cursor_resumen.fetchone()
        cursor_resumen.close()
        
    elif tipo == "mensual":
        cursor.execute('''
            SELECT DATE(timestamp) as fecha, 
                   COALESCE(SUM(personas), 0) as personas,
                   COALESCE(SUM(carros), 0) as carros,
                   COALESCE(SUM(motos), 0) as motos,
                   COALESCE(SUM(buses), 0) as buses,
                   COALESCE(SUM(camiones), 0) as camiones,
                   COALESCE(SUM(total), 0) as total
            FROM detections_log
            WHERE DATE(timestamp) >= CURRENT_DATE - INTERVAL '30 days'
            GROUP BY DATE(timestamp)
            ORDER BY fecha
        ''')
        titulo = f"Reporte Mensual - {datetime.now().strftime('%Y-%m-%d')}"
        nombre = f"reporte_mensual_{datetime.now().strftime('%Y-%m-%d')}.xlsx"
        headers = ['Fecha', 'Personas', 'Carros', 'Motos', 'Buses', 'Camiones', 'Total']
        
        cursor_resumen = conn.cursor()
        cursor_resumen.execute('''
            SELECT 
                COALESCE(SUM(personas), 0) as total_personas,
                COALESCE(SUM(carros), 0) as total_carros,
                COALESCE(SUM(motos), 0) as total_motos,
                COALESCE(SUM(buses), 0) as total_buses,
                COALESCE(SUM(camiones), 0) as total_camiones,
                COALESCE(SUM(total), 0) as total_vehiculos,
                COUNT(*) as total_registros,
                COALESCE(AVG(total), 0) as promedio
            FROM detections_log
            WHERE DATE(timestamp) >= CURRENT_DATE - INTERVAL '30 days'
        ''')
        resumen = cursor_resumen.fetchone()
        cursor_resumen.close()
    
    else:
        cursor.close()
        conn.close()
        return {"error": "Tipo no valido. Usa: diario, semanal o mensual"}
    
    results = cursor.fetchall()
    cursor.close()
    conn.close()
    
    if not results:
        return {"error": "No hay datos para exportar"}
    
    wb = openpyxl.Workbook()
    
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="1a3a5c", end_color="1a3a5c", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    
    thin_border = Border(
        left=Side(style='thin', color='cccccc'),
        right=Side(style='thin', color='cccccc'),
        top=Side(style='thin', color='cccccc'),
        bottom=Side(style='thin', color='cccccc')
    )
    
    ws_resumen = wb.active
    ws_resumen.title = "Resumen"
    
    ws_resumen['A1'] = titulo
    ws_resumen['A1'].font = Font(bold=True, size=16, color="1a3a5c")
    ws_resumen['A1'].alignment = Alignment(horizontal="center", vertical="center")
    ws_resumen.merge_cells('A1:C1')
    
    ws_resumen['A2'] = f"Generado: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
    ws_resumen['A2'].font = Font(size=10, color="666666")
    ws_resumen['A2'].alignment = Alignment(horizontal="center", vertical="center")
    ws_resumen.merge_cells('A2:C2')
    
    ws_resumen['A4'] = "Metrica"
    ws_resumen['A4'].font = header_font
    ws_resumen['A4'].fill = header_fill
    ws_resumen['A4'].alignment = header_alignment
    ws_resumen['A4'].border = thin_border
    
    ws_resumen['B4'] = "Valor"
    ws_resumen['B4'].font = header_font
    ws_resumen['B4'].fill = header_fill
    ws_resumen['B4'].alignment = header_alignment
    ws_resumen['B4'].border = thin_border
    
    ws_resumen['C4'] = "Porcentaje"
    ws_resumen['C4'].font = header_font
    ws_resumen['C4'].fill = header_fill
    ws_resumen['C4'].alignment = header_alignment
    ws_resumen['C4'].border = thin_border
    
    if resumen:
        total_vehiculos = resumen[5] or 0
        total_personas = resumen[0] or 0
        total_carros = resumen[1] or 0
        total_motos = resumen[2] or 0
        total_buses = resumen[3] or 0
        total_camiones = resumen[4] or 0
        total_registros = resumen[6] or 0
        promedio = round(resumen[7], 2) if resumen[7] else 0
    else:
        total_vehiculos = total_personas = total_carros = total_motos = total_buses = total_camiones = total_registros = 0
        promedio = 0
    
    resumen_data = [
        ["Total Vehiculos", total_vehiculos, "100%"],
        ["Total Personas", total_personas, f"{round((total_personas/total_vehiculos)*100, 1) if total_vehiculos > 0 else 0}%"],
        ["Total Carros", total_carros, f"{round((total_carros/total_vehiculos)*100, 1) if total_vehiculos > 0 else 0}%"],
        ["Total Motos", total_motos, f"{round((total_motos/total_vehiculos)*100, 1) if total_vehiculos > 0 else 0}%"],
        ["Total Buses", total_buses, f"{round((total_buses/total_vehiculos)*100, 1) if total_vehiculos > 0 else 0}%"],
        ["Total Camiones", total_camiones, f"{round((total_camiones/total_vehiculos)*100, 1) if total_vehiculos > 0 else 0}%"],
        ["Total Registros", total_registros, ""],
        ["Promedio por Minuto", promedio, ""],
    ]
    
    for i, (metric, value, percentage) in enumerate(resumen_data, 5):
        ws_resumen[f'A{i}'] = metric
        ws_resumen[f'A{i}'].font = Font(size=11)
        ws_resumen[f'A{i}'].alignment = Alignment(horizontal="left", vertical="center")
        ws_resumen[f'A{i}'].border = thin_border
        
        ws_resumen[f'B{i}'] = value
        ws_resumen[f'B{i}'].font = Font(bold=True, size=11, color="1a3a5c")
        ws_resumen[f'B{i}'].alignment = Alignment(horizontal="center", vertical="center")
        ws_resumen[f'B{i}'].border = thin_border
        
        ws_resumen[f'C{i}'] = percentage
        ws_resumen[f'C{i}'].font = Font(size=10, color="666666")
        ws_resumen[f'C{i}'].alignment = Alignment(horizontal="center", vertical="center")
        ws_resumen[f'C{i}'].border = thin_border
    
    ws_resumen.column_dimensions['A'].width = 25
    ws_resumen.column_dimensions['B'].width = 15
    ws_resumen.column_dimensions['C'].width = 15
    
    ws_datos = wb.create_sheet("Datos Detallados")
    
    for col, header in enumerate(headers, 1):
        cell = ws_datos.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border
    
    for row_idx, row in enumerate(results, 2):
        for col_idx, value in enumerate(row, 1):
            cell = ws_datos.cell(row=row_idx, column=col_idx, value=value)
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center", vertical="center")
    
    for col in range(1, len(headers) + 1):
        ws_datos.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 15
    
    with tempfile.NamedTemporaryFile(delete=False, suffix='.xlsx') as tmp:
        wb.save(tmp.name)
        tmp_path = tmp.name
    
    return FileResponse(
        path=tmp_path,
        filename=nombre,
        media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )

if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=8000)